# Copyright (c) Huawei Technologies Co., Ltd. 2026. All rights reserved.

"""write_file / edit_file / read_file Office document support (.docx only)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document

from openjiuwen.harness.tools import filesystem as fs_mod
from openjiuwen.harness.tools.filesystem import EditFileTool, ReadFileTool, WriteFileTool


@pytest.fixture(autouse=True)
def _clear_read_registry() -> None:
    fs_mod._FILE_READ_REGISTRY.clear()
    yield
    fs_mod._FILE_READ_REGISTRY.clear()


def _make_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Write_file")
    doc.add_paragraph("Read_file")
    doc.save(str(path))
    return path


@pytest.mark.asyncio
async def test_write_file_rewrites_docx_after_successful_read(tmp_path: Path) -> None:
    docx_path = _make_docx(tmp_path / "123.docx")
    op = MagicMock()

    read_result = await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    assert read_result.success is True
    assert "Write_file" in (read_result.data or {}).get("content", "")

    write_result = await WriteFileTool(op, language="en").invoke(
        {"file_path": str(docx_path), "content": "Write_file\nRead_file\n你好"}
    )
    assert write_result.success is True, write_result.error
    assert docx_path.read_bytes()[:2] == b"PK"

    verify = await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    assert verify.success is True
    content = (verify.data or {}).get("content", "")
    assert "你好" in content
    assert "Write_file" in content
    assert "Read_file" in content


@pytest.mark.asyncio
async def test_edit_file_updates_docx_via_text_replace(tmp_path: Path) -> None:
    docx_path = _make_docx(tmp_path / "note.docx")
    op = MagicMock()

    await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    edit_result = await EditFileTool(op, language="en").invoke(
        {
            "file_path": str(docx_path),
            "old_string": "Read_file",
            "new_string": "Read_file\n你好",
        }
    )
    assert edit_result.success is True, edit_result.error
    assert docx_path.read_bytes()[:2] == b"PK"

    verify = await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    assert "你好" in (verify.data or {}).get("content", "")


@pytest.mark.asyncio
async def test_write_file_still_requires_read_for_plain_text(tmp_path: Path) -> None:
    txt_path = tmp_path / "note.txt"
    txt_path.write_text("hello\n", encoding="utf-8")

    write_result = await WriteFileTool(MagicMock(), language="en").invoke(
        {"file_path": str(txt_path), "content": "hello\nworld\n"}
    )
    assert write_result.success is False
    assert "has not been read yet" in (write_result.error or "").lower()


@pytest.mark.asyncio
async def test_write_file_creates_new_docx_without_prior_read(tmp_path: Path) -> None:
    docx_path = tmp_path / "new.docx"
    write_result = await WriteFileTool(MagicMock(), language="en").invoke(
        {"file_path": str(docx_path), "content": "hello\nworld"}
    )
    assert write_result.success is True, write_result.error
    assert docx_path.exists()
    assert docx_path.read_bytes()[:2] == b"PK"

    verify = await ReadFileTool(MagicMock(), language="en").invoke({"file_path": str(docx_path)})
    assert "hello" in (verify.data or {}).get("content", "")
    assert "world" in (verify.data or {}).get("content", "")


@pytest.mark.asyncio
async def test_write_file_rewrites_empty_docx_after_read(tmp_path: Path) -> None:
    docx_path = tmp_path / "empty.docx"
    doc = Document()
    doc.save(str(docx_path))
    op = MagicMock()

    read_result = await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    assert read_result.success is True

    write_result = await WriteFileTool(op, language="en").invoke(
        {"file_path": str(docx_path), "content": "hello\nworld"}
    )
    assert write_result.success is True, write_result.error
    assert docx_path.exists()

    verify = await ReadFileTool(op, language="en").invoke({"file_path": str(docx_path)})
    assert "hello" in (verify.data or {}).get("content", "")
    assert "world" in (verify.data or {}).get("content", "")


@pytest.mark.asyncio
async def test_read_file_rejects_legacy_doc(tmp_path: Path) -> None:
    doc_path = tmp_path / "legacy.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"not-a-real-doc")

    result = await ReadFileTool(MagicMock(), language="en").invoke({"file_path": str(doc_path)})
    assert result.success is False
    err = (result.error or "").lower()
    assert "legacy office format" in err
    assert ".doc" in err
    assert "docx" in err or ".docx" in err or "modern format" in err


@pytest.mark.asyncio
async def test_write_file_rejects_legacy_doc(tmp_path: Path) -> None:
    result = await WriteFileTool(MagicMock(), language="en").invoke(
        {"file_path": str(tmp_path / "legacy.doc"), "content": "hello"}
    )
    assert result.success is False
    err = (result.error or "").lower()
    assert ".doc" in err
    assert "convert" in err


@pytest.mark.asyncio
async def test_write_file_docx_skips_blank_lines_from_read_shape(tmp_path: Path) -> None:
    """read_file joins paragraphs with \\n\\n; write must not materialize empty paragraphs."""
    docx_path = tmp_path / "blank.docx"
    write_result = await WriteFileTool(MagicMock(), language="en").invoke(
        {"file_path": str(docx_path), "content": "Write_file\n\nRead_file\n\n你好"}
    )
    assert write_result.success is True, write_result.error

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Write_file", "Read_file", "你好"]
    assert not list(docx_path.parent.glob("*.writing"))


@pytest.mark.asyncio
async def test_write_office_document_keeps_original_when_save_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docx_path = _make_docx(tmp_path / "keep.docx")
    original = docx_path.read_bytes()

    def boom(self, path):  # noqa: ANN001
        raise RuntimeError("simulated save failure")

    monkeypatch.setattr("docx.document.Document.save", boom)

    with pytest.raises(RuntimeError, match="simulated save failure"):
        fs_mod._write_office_document(str(docx_path), "new content")

    assert docx_path.read_bytes() == original
    assert not list(docx_path.parent.glob("*.writing"))


def test_office_text_lines_drops_blank_separators() -> None:
    assert fs_mod._office_text_lines("a\n\nb\n  \nc\n") == ["a", "b", "c"]
    assert fs_mod._office_text_lines("") == []
